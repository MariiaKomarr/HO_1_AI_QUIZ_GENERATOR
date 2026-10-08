from pathlib import Path

import pymupdf
from docx import Document


def read_document(file_path: str) -> str:
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(
            f"Файл не знайдено: {file_path}"
        )

    extension = path.suffix.lower()

    if extension == ".txt":
        text = path.read_text(encoding="utf-8-sig")

    elif extension == ".pdf":
        with pymupdf.open(path) as pdf:
            text = "\n".join(
                page.get_text() for page in pdf
            )

    elif extension == ".docx":
        document = Document(path)
        text = "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
        )

        # Додаємо текст із таблиць
        for table in document.tables:
            for row in table.rows:
                text += "\n" + " | ".join(
                    cell.text for cell in row.cells
                )

    else:
        raise ValueError(
            "Підтримуються лише PDF, DOCX і TXT."
        )

    if not text.strip():
        raise ValueError(
            "У документі не знайдено тексту. "
            "Можливо, це сканований PDF."
        )

    return text.strip()